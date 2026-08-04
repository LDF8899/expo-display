# -*- coding: utf-8 -*-
"""dump docx 段落样式与文本，用于设计章节拆分逻辑。"""
import sys
from pathlib import Path
from docx import Document


def dump(path: Path, limit=120):
    print(f"== {path.name} ==")
    doc = Document(str(path))
    n = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        style = p.style.name if p.style else ""
        if not text:
            continue
        n += 1
        if n > limit:
            print("  ...(truncated)")
            break
        print(f"[{style}] {text[:60]}")
    # 表格信息
    print(f"  -- tables: {len(doc.tables)}, paragraphs total: {len(doc.paragraphs)}")


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        dump(Path(arg))
